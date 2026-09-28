from __future__ import annotations

import re
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

from etl.core.logic.classificacoes_listagens import (
    SITUACOES_NTE,
    ser_de_cota_ingresso,
    ser_de_cota_matricula,
)
from etl.core.logic.models import EstudanteListagem, ListagensExtraidas
from etl.core.logic.normalizacao_nomes import normalizar_nome
from etl.core.logic.resolvers.campus_resolver import normalizar_slug

# Contrato de entrada verificado nos 6 arquivos (spec FR-002).
COLUNAS_CONTRATO = (
    "Matrícula",
    "Nome",
    "Curso",
    "Situação Matrícula",
    "Sexo",
    "Nascimento",
    "Desc_Forma_Ingresso_Matricula",
    "Desc_Cota",
)

PADRAO_ARQUIVO = re.compile(r"^listagem_(\d{4})_([12])\.xlsx$")
RE_TITULO_CAMPUS = re.compile(r"^Campus\s+(.+?)\s*(?:–|-|—)", re.UNICODE)
RE_SEMESTRE_TITULO = re.compile(r"(\d{4})/([12])")
SEMESTRES_ESPERADOS = frozenset({1, 2})


def _normalizar_matricula(valor: object) -> str:
    """Converte a célula de matrícula em chave determinística de str."""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _texto(valor: object) -> str | None:
    """Texto da célula sem espaços, ou None para vazio."""
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


class ListagensXlsxSource:
    """Fonte de extração das planilhas `data/raw/listagem_<AAAA>_<S>.xlsx`."""

    def __init__(self, pasta: Path | str) -> None:
        self.pasta = Path(pasta)

    def extract(self) -> ListagensExtraidas:
        extraidas = ListagensExtraidas()
        arquivos = self._buscar_arquivos()
        semestres_por_chave: dict[tuple[str, int], set[int]] = {}
        campus_titulo_por_ano: dict[int, dict[str, str]] = {}

        for arquivo in arquivos:
            ano, semestre = self._ano_semestre_do_nome(arquivo)
            extraidas.anos_encontrados.add(ano)
            self._ler_arquivo(
                arquivo,
                ano,
                semestre,
                extraidas,
                semestres_por_chave,
                campus_titulo_por_ano,
            )

        self._avisar_semestres_ausentes(extraidas, semestres_por_chave)
        return extraidas

    def _buscar_arquivos(self) -> list[Path]:
        arquivos = []
        for caminho in sorted(self.pasta.glob("listagem_*.xlsx")):
            if PADRAO_ARQUIVO.match(caminho.name):
                arquivos.append(caminho)
        if not arquivos:
            raise ValueError(
                f"nenhum arquivo de listagem encontrado em {self.pasta} "
                f"(esperado listagem_<AAAA>_<1|2>.xlsx)"
            )
        # Ordena por (ano, semestre) para determinismo.
        return sorted(
            arquivos,
            key=lambda a: (
                int(PADRAO_ARQUIVO.match(a.name).group(1)),
                int(PADRAO_ARQUIVO.match(a.name).group(2)),
            ),
        )

    def _ano_semestre_do_nome(self, arquivo: Path) -> tuple[int, int]:
        match = PADRAO_ARQUIVO.match(arquivo.name)
        if not match:  # pragma: no cover — filtrado em _buscar_arquivos
            raise ValueError(
                f"{arquivo.name}: nome fora do contrato listagem_<AAAA>_<S>.xlsx"
            )
        return int(match.group(1)), int(match.group(2))

    def _ler_arquivo(
        self,
        arquivo: Path,
        ano: int,
        semestre: int,
        extraidas: ListagensExtraidas,
        semestres_por_chave: dict[tuple[str, int], set[int]],
        campus_titulo_por_ano: dict[int, dict[str, str]],
    ) -> None:
        try:
            wb = load_workbook(arquivo, read_only=True, data_only=True)
        except (
            Exception
        ) as exc:  # noqa: BLE001 - InvalidFileException/BadZipFile/OSError
            raise ValueError(
                f"{arquivo.name}: falha ao ler arquivo XLSX (corrompido?): {exc}"
            ) from exc

        try:
            ws = wb.active
            linhas = ws.iter_rows(values_only=True)
            titulo = self._ler_titulo(arquivo, linhas)
            self._ler_cabecalho(arquivo, linhas)

            campus_nome = extrair_campus_do_titulo(titulo)
            if campus_nome is None:
                raise ValueError(
                    f"{arquivo.name}: título não identifica o campus — '{titulo}'"
                )
            slug = normalizar_slug(campus_nome)
            extraidas.nomes_campus.setdefault(slug, campus_nome)

            self._conferir_titulo_com_arquivo(
                arquivo, ano, semestre, titulo, extraidas.avisos
            )

            chave = (slug, ano)
            por_semestre = extraidas.por_campus_ano.setdefault(chave, {})
            alunos = por_semestre.setdefault(semestre, [])
            semestres_por_chave.setdefault(chave, set()).add(semestre)

            # Divergência de campus declarado entre semestres do mesmo ano (FR-003).
            nomes_ano = campus_titulo_por_ano.setdefault(ano, {})
            if nomes_ano and slug not in nomes_ano:
                extraidas.avisos.append(
                    f"campus divergente entre semestres do ano {ano}: "
                    f"'{nomes_ano[next(iter(nomes_ano))]}' vs '{campus_nome}'"
                )
            nomes_ano[slug] = campus_nome

            self._ler_alunos(
                arquivo,
                linhas,
                alunos,
                extraidas.nomes_cotistas,
                chave,
                extraidas.avisos,
            )
        finally:
            wb.close()

    def _ler_titulo(self, arquivo: Path, linhas: Iterator[tuple]) -> str:
        try:
            primeira = next(linhas)
        except StopIteration as exc:
            raise ValueError(f"{arquivo.name}: planilha vazia") from exc
        titulo = _texto(primeira[0] if primeira else None)
        if titulo is None:
            raise ValueError(f"{arquivo.name}: linha 1 (título com campus) ausente")
        return titulo

    def _ler_cabecalho(self, arquivo: Path, linhas: Iterator[tuple]) -> tuple[str, ...]:
        try:
            next(linhas)  # linha 2 (vazia no contrato)
            linha_header = next(linhas)
        except StopIteration as exc:
            raise ValueError(
                f"{arquivo.name}: cabeçalho ausente (esperado na linha 3)"
            ) from exc
        cab = tuple(_texto(v) for v in linha_header)
        if cab != COLUNAS_CONTRATO:
            raise ValueError(
                f"{arquivo.name}: cabeçalho fora do contrato de 8 colunas "
                f"(esperado linha 3: {', '.join(COLUNAS_CONTRATO)})"
            )
        return cab

    def _conferir_titulo_com_arquivo(
        self,
        arquivo: Path,
        ano: int,
        semestre: int,
        titulo: str,
        avisos: list[str],
    ) -> None:
        match = RE_SEMESTRE_TITULO.search(titulo)
        if match is None:
            avisos.append(
                f"{arquivo.name}: título sem indicação de semestre ('Semestres letivo: AAAA/S')"
            )
            return
        ano_titulo, sem_titulo = int(match.group(1)), int(match.group(2))
        if (ano_titulo, sem_titulo) != (ano, semestre):
            avisos.append(
                f"{arquivo.name}: título indica {ano_titulo}/{sem_titulo}, "
                f"mas o nome do arquivo indica {ano}/{semestre}"
            )

    def _ler_alunos(
        self,
        arquivo: Path,
        linhas: Iterator[tuple],
        alunos: list[EstudanteListagem],
        nomes_cotistas: dict[tuple[str, int], set[str]],
        chave: tuple[str, int],
        avisos: list[str],
    ) -> None:
        for n_linha, linha in enumerate(linhas, start=4):
            if not linha or all(v is None for v in linha):
                continue
            matricula_raw = linha[0]
            if matricula_raw is None or _texto(matricula_raw) is None:
                avisos.append(
                    f"{arquivo.name}: linha {n_linha} ignorada — matrícula ausente"
                )
                continue

            estudante = EstudanteListagem(
                matricula=_normalizar_matricula(matricula_raw),
                situacao=_texto(linha[3] if len(linha) > 3 else None),
                forma_ingresso=_texto(linha[6] if len(linha) > 6 else None),
                forma_matricula_cota=_texto(linha[7] if len(linha) > 7 else None),
            )
            alunos.append(estudante)

            # Retém apenas nomes de cotistas (NTE + ambas colunas de cota), já
            # normalizados, para o cruzamento NTECPP (FR-012). Em memória apenas.
            if (
                estudante.situacao in SITUACOES_NTE
                and ser_de_cota_ingresso(estudante.forma_ingresso)
                and ser_de_cota_matricula(estudante.forma_matricula_cota)
            ):
                nome = _texto(linha[1] if len(linha) > 1 else None)
                normalizado = normalizar_nome(nome)
                if normalizado:
                    nomes_cotistas.setdefault(chave, set()).add(normalizado)

    def _avisar_semestres_ausentes(
        self,
        extraidas: ListagensExtraidas,
        semestres_por_chave: dict[tuple[str, int], set[int]],
    ) -> None:
        for chave, encontrados in sorted(semestres_por_chave.items()):
            slug, ano = chave
            for semestre in sorted(SEMESTRES_ESPERADOS - encontrados):
                extraidas.avisos.append(
                    f"listagem_{ano}_{semestre}.xlsx não encontrado — "
                    f"semestre {semestre} ausente para {slug}/{ano}"
                )


def extrair_campus_do_titulo(titulo: str) -> str | None:
    match = RE_TITULO_CAMPUS.match(titulo.strip())
    if not match:
        return None
    return match.group(1).strip()
